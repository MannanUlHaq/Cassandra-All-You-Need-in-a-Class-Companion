import modal
import os
import json
import re
import uuid
import tempfile
import shutil
import hashlib
import time
from pathlib import Path
from io import BytesIO
from typing import Dict, List, Optional
from datetime import datetime

# --- 1. Define the Modal Environment ---
image = modal.Image.debian_slim(python_version="3.10").apt_install(
    "libgl1", "libglib2.0-0", "build-essential", "git" # Required for OpenCV and cloning Unsloth
).pip_install(
    "python-pptx", "pymupdf", "lxml", "opencv-python-headless", "pillow",
    "torch", "torchvision", "transformers", "accelerate", "huggingface_hub", 
    "fastapi[standard]", "python-multipart", "numpy", "openai" # Added openai for DeepSeek API
).run_commands(
    'pip install "unsloth[colab-new] @ git+https://github.com/unslothai/unsloth.git"'
)

app = modal.App("slides-ai-backend")

# Create a persistent volume to store your finetuned models!
model_volume = modal.Volume.from_name("finetuned-models-vol", create_if_missing=True)

# ---------------- CONFIG / CONSTANTS ----------------
CLASSIFIER_MODEL_PATH = "/models/table-equation-classifier"
CAPTION_MODEL_PATH = "/models/blip2_captioning_model"
QWEN_MODEL_PATH = "/models/qwen_fine_tuned_final"
DEEPSEEK_MODEL_ID = "deepseek-chat" # Updated to DeepSeek's official chat model ID

CAPTION_PROMPT = "A photo of" 

NS = {
    'p':  'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a':  'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r':  'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'm':  'http://schemas.openxmlformats.org/officeDocument/2006/math',
    'pic':'http://schemas.openxmlformats.org/drawingml/2006/picture'
}

PROMPT_TABLE = """You are a specialist in recognizing and generating LaTeX tables.
1. Output ONLY the LaTeX code for the table in the image.
2. Do not include any introductory text or explanations.
3. Ensure the column and row structure is preserved exactly.
4. If the table lines are invisible, infer the structure based on alignment."""

PROMPT_EQUATION = """You are a mathematical OCR assistant.
1. Transcribe the content of the image into LaTeX.
2. The image may contain a single equation, multiple equations, or a mix of English text and math.
3. Preserve the exact order of text and math as shown.
4. Use standard LaTeX formatting for all mathematical symbols.
5. Do not include any introductory text or explanations."""

# ==============================================================================
# ---------------- PIPELINE HELPERS (PRESERVING EXACT LOGIC) -------------------
# ==============================================================================

def render_pdf_to_images(pdf_path, outdir, zoom=2.0, out_prefix="slide_"):
    import fitz  # PyMuPDF
    doc = fitz.open(pdf_path)
    saved = []
    for i, page in enumerate(doc):
        try:
            mat = fitz.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            out_path = outdir / f"{out_prefix}{i+1}.png"
            pix.save(out_path)
            saved.append(str(out_path))
        except Exception as e:
            print(f"⚠️ Error rendering page {i+1}: {e}")
            continue
    doc.close()
    return saved

def get_local_xfrm(elem):
    xfrm = elem.find('.//a:xfrm', namespaces=NS) or elem.find('.//p:xfrm', namespaces=NS)
    if xfrm is None: return None
    off = xfrm.find('a:off', namespaces=NS); ext = xfrm.find('a:ext', namespaces=NS)
    if off is None or ext is None: return None
    try: return (int(off.get('x',0)), int(off.get('y',0)), int(ext.get('cx',0)), int(ext.get('cy',0)))
    except Exception: return None

def xml_has_table(elem):
    if elem.find('.//a:tbl', namespaces=NS) is not None: return True
    gd = elem.find('.//a:graphicData', namespaces=NS)
    if gd is not None and 'table' in (gd.get('uri','') or '').lower(): return True
    return False

def xml_has_equation(elem):
    if elem.find('.//m:oMath', namespaces=NS) is not None or elem.find('.//m:oMathPara', namespaces=NS) is not None: return True
    ole = elem.find('.//p:oleObj', namespaces=NS)
    if ole is not None and any(k in (ole.get('progId','') or '').lower() for k in ['equation','math','mathtype','omml']): return True
    return False

def xml_has_text(elem):
    for t in elem.findall('.//a:t', namespaces=NS):
        if t.text and t.text.strip(): return True
    return False

def read_text_from_elem(elem):
    texts = [t.text for t in elem.findall('.//a:t', namespaces=NS) if t.text]
    if not texts: return None
    return "\n".join(texts).strip()

def md5_of_image(img):
    bio = BytesIO()
    img.save(bio, format="PNG")
    return hashlib.md5(bio.getvalue()).hexdigest()

def save_pil_image(img, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, format="PNG")
    return str(path)

def are_images_identical(path1, path2):
    import numpy as np
    from PIL import Image
    try:
        img1 = Image.open(path1).convert("RGB")
        img2 = Image.open(path2).convert("RGB")
        if img1.size != img2.size: return False
        return np.array_equal(np.array(img1), np.array(img2))
    except Exception: return False

def emu_to_pixel_rect(emu_bbox, slide_emu_w, slide_emu_h, img_w, img_h):
    x, y, cx, cy = emu_bbox
    sx = img_w / float(slide_emu_w)
    sy = img_h / float(slide_emu_h)
    left  = int(round(x * sx))
    top   = int(round(y * sy))
    right = int(round((x + cx) * sx))
    bottom= int(round((y + cy) * sy))
    left, top = max(0, left), max(0, top)
    right, bottom = min(img_w, right), min(img_h, bottom)
    if right <= left: right = min(img_w, left + 1)
    if bottom <= top: bottom = min(img_h, top + 1)
    return left, top, right, bottom

def obtain_image_for_elem(elem, slide_obj, zipf, slide_index, e_idx, media_dir):
    if slide_obj is not None:
        for shp in slide_obj.shapes:
            try:
                if shp.element is elem or (shp.element is not None and id(shp.element) == id(elem)):
                    if getattr(shp, 'image', None) is not None:
                        img = shp.image
                        fpath = media_dir / f"slide{slide_index}_e{e_idx}.{img.ext}"
                        with open(fpath, 'wb') as f: f.write(img.blob)
                        return str(fpath)
            except Exception: continue
    try:
        blip = elem.find('.//a:blip', namespaces=NS)
        if blip is not None:
            rid = blip.get('{%s}embed' % NS['r'])
            if rid and slide_obj is not None:
                rel = slide_obj.part.rels.get(rid)
                if rel is not None:
                    target = rel.target_ref
                    if target.startswith('../'): target = target[3:]
                    data = zipf.read(target)
                    ext = Path(target).suffix.lstrip('.')
                    fpath = media_dir / f"slide{slide_index}_e{e_idx}.{ext}"
                    with open(fpath, 'wb') as f: f.write(data)
                    return str(fpath)
    except Exception: pass
    return None

def split_equation_image(image_path):
    import cv2
    img = cv2.imread(image_path)
    if img is None: return []
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    _, thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (30, 10))
    dilated = cv2.dilate(thresh, kernel, iterations=1)
    hist = cv2.reduce(dilated, 1, cv2.REDUCE_AVG).reshape(-1)

    h, w = img.shape[:2]
    raw_lines, in_line, start_y, threshold = [], False, 0, 1
    for y in range(h):
        val = hist[y]
        if val > threshold and not in_line:
            in_line = True; start_y = y
        elif val <= threshold and in_line:
            in_line = False; end_y = y
            if (end_y - start_y) > 5: raw_lines.append((start_y, end_y))
    if in_line: raw_lines.append((start_y, h))

    merged_lines = []
    if raw_lines:
        curr_y1, curr_y2 = raw_lines[0]
        for i in range(1, len(raw_lines)):
            next_y1, next_y2 = raw_lines[i]
            if (next_y1 - curr_y2) < 30: curr_y2 = next_y2
            else:
                merged_lines.append((curr_y1, curr_y2))
                curr_y1, curr_y2 = next_y1, next_y2
        merged_lines.append((curr_y1, curr_y2))

    results = []
    for y1, y2 in merged_lines:
        y1_pad = max(0, y1 - 5)
        y2_pad = min(h, y2 + 5)
        results.append((y1_pad, y2_pad, img[y1_pad:y2_pad, 0:w]))
    return results

def classify_single_image(image_path, model, processor, device):
    import torch
    from PIL import Image
    try:
        image = Image.open(image_path).convert("RGB")
        inputs = processor(images=image, return_tensors="pt").to(device)
        with torch.no_grad(): outputs = model(**inputs)
        predicted_class_idx = outputs.logits.argmax(-1).item()
        return model.config.id2label[predicted_class_idx]
    except Exception as e:
        print(f"  Error classifying {image_path}: {e}")
        return None

def generate_image_caption(image_path, model, processor, device):
    import torch
    from PIL import Image
    try:
        image = Image.open(image_path).convert("RGB")
        dtype = torch.float16 if device == "cuda" else torch.float32
        inputs = processor(images=image, text=CAPTION_PROMPT, return_tensors="pt").to(device, dtype)
        with torch.no_grad():
            generated_ids = model.generate(**inputs, max_new_tokens=50)
        return processor.batch_decode(generated_ids, skip_special_tokens=True)[0].strip()
    except Exception as e:
        print(f"  Error captioning {image_path}: {e}")
        return None

# ==============================================================================
# ------------------ CLASS FOR DEEPSEEK RESTRUCTURING --------------------------
# ==============================================================================

class SlideStructureProcessor:
    def __init__(self, api_key: str, model: str):
        from openai import OpenAI
        # DeepSeek API uses OpenAI compatible SDK
        self.client = OpenAI(api_key=api_key, base_url="https://api.deepseek.com")
        self.model = model
        self.std_width = 1920
        self.std_height = 1080

    def normalize_bbox(self, bbox: Optional[List[float]]) -> Dict[str, int]:
        if not bbox or len(bbox) != 4: return {'y': 0, 'x': 0, 'w': 0, 'h': 0}
        l, t, r, b = bbox
        return {
            'y': int((t / self.std_height) * 1000), 'x': int((l / self.std_width) * 1000),
            'w': int((max(0, r - l) / self.std_width) * 1000), 'h': int((max(0, b - t) / self.std_height) * 1000)
        }

    def _get_element_context(self, el: Dict) -> Dict:
        norm = self.normalize_bbox(el.get('px_bbox_estimate'))
        elem_type = el.get('type', 'unknown')
        if elem_type == 'text': content_snippet = (el.get('text') or "").strip()
        elif elem_type == 'image': content_snippet = el.get('caption') or "No caption provided"
        elif elem_type in ['equation', 'table']: content_snippet = el.get('extracted_data') or ""
        else: content_snippet = ""
            
        context = {"id": el['id'], "type": elem_type, "content": content_snippet, "spatial": f"top={norm['y']}, left={norm['x']}, w={norm['w']}, h={norm['h']}"}
        if "image_path" in el: context["image_path"] = el["image_path"]
        return context

    def process_slide_structure(self, slide_data: Dict) -> Dict:
        slide_idx = slide_data.get('slide_index')
        elements = slide_data.get('elements', [])
        if not elements: return {"slide_index": slide_idx, "structured_elements": []}

        elements_context = [self._get_element_context(el) for el in elements]
        elements_context.sort(key=lambda x: int(x['spatial'].split(',')[0].split('=')[1]))

        system_prompt = """You are an advanced Document Layout Analysis AI.
Your task is to reconstruct the correct logical reading order of elements extracted from a PowerPoint slide by using a deep understanding of the element content together with spatial layout information.
The extracted elements already contain finalized content:
- Text elements contain their full textual content
- Image elements contain their captions or descriptions in the content field
- Equation elements contain their equation content
- Table elements contain their table content
You MUST NOT modify, rewrite, summarize, merge, or correct any content.
Your task is ONLY to determine the correct logical placement and reading order of the elements.
Raw slide extraction may be incorrect and can contain:
- Incorrect reading order
- Visually separated but conceptually related elements
- Semantic imbalance (e.g., headings appearing after definitions)
You must fully understand the meaning of each element and determine how a human would naturally read the slide.
If an image and a text element are conceptually related (e.g., an image and its caption), they should appear consecutively in the correct logical order, but their content must remain unchanged.
INPUT FORMAT:
Each slide element is provided as a JSON object with the following fields:
- id: unique element identifier
- type: one of [text, image, equation, table]
- content: extracted textual or symbolic content (DO NOT MODIFY)
- spatial: spatial position in the form "top=..., left=..., w=..., h=..."
- image_path: (Optional) original file path for the element
OUTPUT REQUIREMENTS:
Return VALID JSON ONLY with the following structure:
{
  "ordered_elements": [
    {
      "id": "...",
      "type": "text | image | equation | table",
      "content": "...",
      "image_path": "... (only if present in input)",
      "rationale": "Brief explanation of why this element is placed at this position"
    }
  ]
}
IMPORTANT CONSTRAINTS:
- NEVER change the content field
- NEVER add new elements
- NEVER remove elements
- Do not add explanations outside JSON
- Ensure the output is valid JSON
"""
        prompt_messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": json.dumps(elements_context, indent=2)}
        ]

        for attempt in range(3):
            try:
                # Updated to use official OpenAI API formatting
                response = self.client.chat.completions.create(
                    model=self.model, 
                    messages=prompt_messages, 
                    max_tokens=8000, 
                    temperature=0.1, 
                    top_p=0.9, 
                    stream=False
                )
                content = response.choices[0].message.content
                clean_text = content.strip()[content.find('{'):]
                clean_text = re.sub(r'(?<!\\)\\(?!["\\/bfnrtu])', r'\\\\', clean_text)
                data, _ = json.JSONDecoder().raw_decode(clean_text)
                
                ordered_list = data.get("ordered_elements", [])
                id_map = {el['id']: el for el in elements}
                final_structure, processed_ids = set(), set()
                final_structure = []
                
                for item in ordered_list:
                    orig_id = item.get("id")
                    if orig_id in id_map:
                        processed_ids.add(orig_id)
                        clean_el = {"id": orig_id, "type": item.get("type", "Body Text"), "content": item.get("content", ""), "rationale": item.get("rationale", "")}
                        if item.get("image_path"): clean_el["image_path"] = item.get("image_path")
                        final_structure.append(clean_el)
                
                # Fallback for dropped elements
                for el in elements:
                    if el['id'] not in processed_ids:
                        c = el.get('text', "") if el['type'] == 'text' else f"[Image: {el.get('caption', 'No caption')}]" if el['type'] == 'image' else el.get('extracted_data', "")
                        fallback_el = {"id": el['id'], "type": el['type'], "content": c, "rationale": "Recovered missing element"}
                        if "image_path" in el: fallback_el["image_path"] = el["image_path"]
                        final_structure.append(fallback_el)
                        
                return {"slide_index": slide_idx, "structured_elements": final_structure}
            except Exception as e:
                time.sleep(5) 
        return {"slide_index": slide_idx, "error": "Failed after retries", "structured_elements": elements}

# ==============================================================================
# ------------------------- MODAL GPU EXECUTION --------------------------------
# ==============================================================================

# 1. Renamed 'container_idle_timeout' to 'scaledown_window'
# 2. Moved 'timeout=3600' HERE to the class decorator
@app.cls(gpu="A10G", image=image, volumes={"/models": model_volume}, scaledown_window=300, timeout=3600)
class SlideProcessorGPU:
    @modal.method()  # <-- Removed the timeout from here
    def process(self, pdf_bytes: bytes, pptx_bytes: bytes, deepseek_api_key: str):
        # Import Unsloth FIRST to apply optimizations before transformers loads
        from unsloth import FastVisionModel
        import torch
        from transformers import AutoImageProcessor, AutoModelForImageClassification, Blip2Processor, Blip2ForConditionalGeneration
        
        import zipfile
        from pptx import Presentation
        from lxml import etree as ET
        from PIL import Image
        import cv2

        temp_dir = Path(tempfile.mkdtemp())
        
        # Setup paths matching user logic
        pdf_path = temp_dir / "input.pdf"
        pptx_path = temp_dir / "input.pptx"
        images_dir = temp_dir / "slides_images" / "my_slides"
        outdir = temp_dir / "pptx_extracted"
        media_dir = outdir / "media"
        table_dir = outdir / "tables"
        eq_dir = outdir / "equations"
        
        # Write files
        pdf_path.write_bytes(pdf_bytes)
        pptx_path.write_bytes(pptx_bytes)
        images_dir.mkdir(parents=True, exist_ok=True)
        media_dir.mkdir(parents=True, exist_ok=True)
        table_dir.mkdir(parents=True, exist_ok=True)
        eq_dir.mkdir(parents=True, exist_ok=True)

        try:
            # ==========================================
            # PIPELINE STEP 1: RENDER PDF TO IMAGES
            # ==========================================
            print("--- 1. Rendering PDF -> PNG ---")
            render_pdf_to_images(pdf_path, images_dir, zoom=2.0)

            # ==========================================
            # PIPELINE STEP 2: DECOMPOSE PPTX
            # ==========================================
            print("--- 2. Extracting Elements ---")
            prs = Presentation(pptx_path)
            zipf = zipfile.ZipFile(pptx_path, 'r')
            slide_w_emu = int(prs.slide_width)
            slide_h_emu = int(prs.slide_height)
            partname_to_slide = {s.part.partname.lstrip('/'): s for s in prs.slides}
            slide_entries = sorted([name for name in zipf.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)], key=lambda n: int(re.search(r'(\d+)', n).group(1)))

            results = {"presentation": str(pptx_path), "total_slides": len(slide_entries), "slides": []}
            seen_texts, seen_image_hashes = set(), set()

            for slide_entry in slide_entries:
                xml_bytes = zipf.read(slide_entry)
                root = ET.fromstring(xml_bytes)
                slide_index = int(re.search(r'(\d+)', slide_entry).group(1))
                slide_obj = partname_to_slide.get(slide_entry)
                slide_img_path = images_dir / f"slide_{slide_index}.png"
                slide_img_exists = slide_img_path.exists()
                slide_img = Image.open(slide_img_path) if slide_img_exists else None
                img_w, img_h = slide_img.size if slide_img else (int(slide_w_emu//12700), int(slide_h_emu//12700)) 
                
                elems = root.xpath(".//p:sp | .//p:pic | .//p:graphicFrame | .//p:oleObj | .//p:grpSp", namespaces=NS)
                slide_elements, current_slide_table_paths = [], []

                for e_idx, elem in enumerate(elems):
                    try:
                        emu_bbox = get_local_xfrm(elem)
                        if not emu_bbox: continue

                        tag = ET.QName(elem.tag).localname.lower()
                        element_type = "image" if tag == "pic" else "table" if xml_has_table(elem) else "equation" if xml_has_equation(elem) else "chart" if elem.find('.//a:graphicData', namespaces=NS) is not None and any(k in (elem.find('.//a:graphicData', namespaces=NS).get('uri') or '').lower() for k in ['chart','diagram','smartart']) else "text"

                        elrec = {"id": f"s{slide_index}_{e_idx}", "type": element_type, "emu_bbox": emu_bbox, "px_bbox_estimate": None, "text": None, "image_path": None}
                        try: elrec['px_bbox_estimate'] = list(emu_to_pixel_rect(emu_bbox, slide_w_emu, slide_h_emu, img_w, img_h))
                        except: pass

                        # Extract media/text
                        imgp = obtain_image_for_elem(elem, slide_obj, zipf, slide_index, e_idx, media_dir)
                        if imgp:
                            elrec['image_path'] = imgp
                            try:
                                h = md5_of_image(Image.open(imgp).convert("RGB"))
                                if h in seen_image_hashes: elrec['duplicate'] = True
                                else: seen_image_hashes.add(h)
                            except: pass

                        if element_type == "text" and xml_has_text(elem):
                            text = read_text_from_elem(elem)
                            if text and text.strip():
                                normalized = re.sub(r'\s+', ' ', text).strip()
                                if normalized in seen_texts: elrec['duplicate'] = True
                                else: seen_texts.add(normalized)
                                elrec['text'] = normalized

                        if element_type in ("table", "equation", "chart") and slide_img is not None and elrec['px_bbox_estimate']:
                            l,t,r,b = elrec['px_bbox_estimate']
                            try:
                                crop = slide_img.crop((l,t,r,b))
                                subdir = table_dir if element_type == "table" else eq_dir if element_type == "equation" else media_dir
                                fpath = subdir / f"slide{slide_index}_e{e_idx}_{element_type}.png"
                                save_pil_image(crop, fpath)
                                elrec['image_path'] = str(fpath)

                                if element_type == "table":
                                    is_dup = any(are_images_identical(str(fpath), existing) for existing in current_slide_table_paths)
                                    if is_dup:
                                        elrec['duplicate'] = True
                                        try: os.remove(fpath); elrec['image_path'] = None
                                        except: pass
                                    else: current_slide_table_paths.append(str(fpath))
                                
                                if not elrec.get('duplicate'):
                                    h = md5_of_image(crop)
                                    if h in seen_image_hashes: elrec['duplicate'] = True
                                    else: seen_image_hashes.add(h)
                            except Exception as e: elrec.setdefault('notes', []).append(f"crop_failed:{e}")

                        # Fallback crop
                        if element_type == "image" and elrec['image_path'] is None and slide_img is not None and elrec['px_bbox_estimate']:
                            l,t,r,b = elrec['px_bbox_estimate']
                            try:
                                crop = slide_img.crop((l,t,r,b))
                                fpath = media_dir / f"slide{slide_index}_e{e_idx}_image.png"
                                save_pil_image(crop, fpath)
                                elrec['image_path'] = str(fpath)
                                h = md5_of_image(crop)
                                if h in seen_image_hashes: elrec['duplicate'] = True
                                else: seen_image_hashes.add(h)
                            except Exception as e: elrec.setdefault('notes', []).append(f"image_crop_failed:{e}")

                        if elrec.get('text') or elrec.get('image_path') and not elrec.get('duplicate'):
                            slide_elements.append(elrec)
                    except Exception as e: slide_elements.append({"id": f"s{slide_index}_{e_idx}", "type": "error", "error": repr(e)})
                results['slides'].append({"slide_index": slide_index, "elements": slide_elements})
            zipf.close()

            # --- Classify Media ---
            device = "cuda" if torch.cuda.is_available() else "cpu"
            print("\n--- Running Classification on Extracted Media ---")
            if os.path.exists(CLASSIFIER_MODEL_PATH):
                cls_processor = AutoImageProcessor.from_pretrained(CLASSIFIER_MODEL_PATH)
                cls_model = AutoModelForImageClassification.from_pretrained(CLASSIFIER_MODEL_PATH).to(device).eval()
                for slide in results['slides']:
                    for el in slide['elements']:
                        if el.get('type') == 'image' and el.get('image_path') and not el.get('duplicate'):
                            orig_path = Path(el['image_path'])
                            if orig_path.exists():
                                label = classify_single_image(str(orig_path), cls_model, cls_processor, device)
                                if label:
                                    label_lower = label.lower()
                                    target_dir = table_dir if "table" in label_lower else eq_dir if "equation" in label_lower else None
                                    new_type = "table" if "table" in label_lower else "equation" if "equation" in label_lower else None
                                    if target_dir and new_type:
                                        try:
                                            new_name = f"{orig_path.stem}_{new_type}{orig_path.suffix}"
                                            new_path = target_dir / new_name
                                            shutil.move(str(orig_path), str(new_path))
                                            el['type'] = new_type; el['image_path'] = str(new_path); el.setdefault('notes', []).append(f"ML_classified_as_{new_type}")
                                        except Exception: pass
                del cls_model, cls_processor; torch.cuda.empty_cache()

            # --- Separate Equations ---
            print("\n--- Checking for Multiple Equations ---")
            for slide in results['slides']:
                final_elements = []
                for el in slide['elements']:
                    if el.get('type') == 'equation' and el.get('image_path') and not el.get('duplicate'):
                        eq_path = el['image_path']
                        splits = split_equation_image(eq_path)
                        if len(splits) > 1:
                            original_bbox = el.get('px_bbox_estimate')
                            for i, (y1, y2, crop_img) in enumerate(splits):
                                new_fpath = eq_dir / f"{Path(eq_path).stem}_part{i+1}{Path(eq_path).suffix}"
                                cv2.imwrite(str(new_fpath), crop_img)
                                new_el = el.copy()
                                new_el['id'] = f"{el['id']}_part{i+1}"; new_el['image_path'] = str(new_fpath)
                                if original_bbox: new_el['px_bbox_estimate'] = [original_bbox[0], original_bbox[1] + y1, original_bbox[2], original_bbox[1] + y2]
                                final_elements.append(new_el)
                        else: final_elements.append(el)
                    else: final_elements.append(el)
                slide['elements'] = final_elements

            # --- Caption Remaining Images ---
            print("\n--- Running Captioning on Remaining Images ---")
            if os.path.exists(CAPTION_MODEL_PATH):
                cap_processor = Blip2Processor.from_pretrained(CAPTION_MODEL_PATH)
                cap_model = Blip2ForConditionalGeneration.from_pretrained(CAPTION_MODEL_PATH, torch_dtype=torch.float16).to(device)
                for slide in results['slides']:
                    for el in slide['elements']:
                        if el.get('type') == 'image' and el.get('image_path') and not el.get('duplicate'):
                            cap = generate_image_caption(el['image_path'], cap_model, cap_processor, device)
                            if cap: el['caption'] = cap
                del cap_model, cap_processor; torch.cuda.empty_cache()


            # ==========================================
            # PIPELINE STEP 3: QWEN LATEX EXTRACTION
            # ==========================================
            print("\n--- 3. Qwen Extraction Pipeline ---")
            if os.path.exists(QWEN_MODEL_PATH):
                model, tokenizer = FastVisionModel.from_pretrained(QWEN_MODEL_PATH, load_in_4bit=True, use_gradient_checkpointing="unsloth")
                FastVisionModel.for_inference(model)
                model.to(device)
                
                for slide in results['slides']:
                    for element in slide.get('elements', []):
                        el_type = element.get('type')
                        img_path = element.get('image_path')
                        if el_type in ['table', 'equation'] and img_path and not element.get('duplicate') and os.path.exists(img_path):
                            specific_prompt = PROMPT_TABLE if el_type == 'table' else PROMPT_EQUATION
                            try:
                                img = Image.open(img_path).convert('RGB')
                                messages = [{"role": "user", "content": [{"type": "image"}, {"type": "text", "text": specific_prompt}]}]
                                input_text = tokenizer.apply_chat_template(messages, add_generation_prompt=True)
                                inputs = tokenizer(img, input_text, add_special_tokens=False, return_tensors="pt").to(device)
                                
                                output_ids = model.generate(**inputs, max_new_tokens=1024, temperature=0.1)
                                generated_tokens = output_ids[0, inputs.input_ids.shape[1]:]
                                cleaned_text = tokenizer.decode(generated_tokens, skip_special_tokens=True).replace("<|im_end|>", "").strip()
                                element['extracted_data'] = cleaned_text
                            except Exception as e: element['extraction_error'] = str(e)
                del model, tokenizer; torch.cuda.empty_cache()


            # ==========================================
            # PIPELINE STEP 4: DEEPSEEK STRUCTURING
            # ==========================================
            print("\n--- 4. DeepSeek Structuring ---")
            processor = SlideStructureProcessor(deepseek_api_key, DEEPSEEK_MODEL_ID)
            structured_slides = []
            
            for slide in results['slides']:
                result = processor.process_slide_structure(slide)
                structured_slides.append(result)

            # --- MODIFY PATHS TO BE RELATIVE ---
            for slide in structured_slides:
                for el in slide.get('structured_elements', []):
                    if 'image_path' in el and el['image_path']:
                        try:
                            # Convert absolute temp path to relative path inside the zippable folder
                            rel_path = Path(el['image_path']).relative_to(outdir)
                            # Standardize to forward slash for JSON usability across OS
                            el['image_path'] = str(rel_path).replace("\\", "/")
                        except ValueError:
                            pass

            final_output = {
                # Clean up the presentation name logic to omit temp paths
                "presentation_info": Path(results.get('presentation', "input.pptx")).name,
                "processing_timestamp": datetime.now().isoformat(),
                "slides": structured_slides
            }

            # Save the JSON file alongside the extracted media in the output directory
            json_path = outdir / "result.json"
            with open(json_path, "w") as f:
                json.dump(final_output, f, indent=2)

            # Zip the directory
            print("\n--- 5. Creating ZIP Archive ---")
            archive_path = temp_dir / "slides_output"
            shutil.make_archive(str(archive_path), 'zip', str(outdir))

            # Read ZIP into memory to return
            with open(f"{archive_path}.zip", "rb") as f:
                zip_bytes = f.read()

            return zip_bytes

        finally:
            # Clean up temporary container storage
            shutil.rmtree(temp_dir, ignore_errors=True)

# --- 3. Web API Endpoint ---
@app.function(image=image, timeout=3600)  # <-- ADD TIMEOUT HERE
@modal.asgi_app()
def fastapi_app():
    from fastapi import FastAPI, UploadFile, File, Form, Response
    web_app = FastAPI()

    @web_app.post("/process-slides")
    async def process_slides(
        pdf_file: UploadFile = File(...),   # <--- Now accepts both files!
        pptx_file: UploadFile = File(...),
        deepseek_api_key: str = Form(...)   # Updated to receive DeepSeek API Key
    ):
        pdf_bytes = await pdf_file.read()
        pptx_bytes = await pptx_file.read()
        
        processor = SlideProcessorGPU()
        # Expect zip file bytes returned from the remote process
        zip_bytes = processor.process.remote(pdf_bytes, pptx_bytes, deepseek_api_key)
        
        # Return the folder contents bundled as a ZIP download!
        return Response(
            content=zip_bytes,
            media_type="application/zip",
            headers={"Content-Disposition": "attachment; filename=slides_extracted_data.zip"}
        )

    return web_app